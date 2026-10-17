from pathlib import Path
import re,json,hashlib
from docx import Document
root=Path.cwd(); out=root/'data'/'final';out.mkdir(parents=True,exist_ok=True)
docpath=root/'末日车队_游戏设计最终总稿.docx';doc=Document(docpath)
blocks=[];cur=None
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  if cur:blocks.append(cur)
  cur={'heading':p.text,'style':p.style.name,'lines':[]}
 elif cur and p.text.strip():cur['lines'].append(p.text)
if cur:blocks.append(cur)
cards={};endings={};epilogues={}
for b in blocks:
 h=re.sub(r'^N ([FXC][0-9]{2}) ',r'N-\1 ',b['heading']);m=re.match(r'^(M\d{2}|T\d{2}|V\d{2}|S\d{2}|R\d{2}|C\d{2}|N-[FXC]\d{2})\s+(.+)',h)
 if m and b['style']=='Heading 3' and (m[1][0]!='S' or m[2] in ['队友争执','资源丢失','夜间求救','车坏了','留下一本书'] or '合格' in '\n'.join(b['lines'])[:90]):
  ident,title=m.groups();text='\n\n'.join(b['lines']);meta=b['lines'][0] if b['lines'] else ''
  # Numbered references such as 回①/③ are prose, never extra choices.
  option_text='\n'.join(line for line in b['lines'] if '①' in line or re.match(r'^[②③④⑤⑥]',line))
  boundaries=[]
  for mark in '①②③④⑤⑥':
   hits=list(re.finditer(re.escape(mark),option_text))
   hits=[x for x in hits if not re.search(r'[回/／、至]$',option_text[:x.start()]) and (not boundaries or x.start()>boundaries[-1][1])]
   if not hits:break
   boundaries.append((mark,hits[0].start()))
  opts=[]
  for i,(mark,start) in enumerate(boundaries):
   raw=option_text[start+1:boundaries[i+1][1] if i+1<len(boundaries) else len(option_text)].strip()
   raw=re.split(r'无(?:人)?处理|无人回应|未处理按',raw)[0].strip()
   label=re.split(r'[：:。；]',raw)[0].strip()
   opts.append({'key':str(i),'text':label,'raw':raw,'inquiry':label.startswith(('追问','问','看看','怎么','为什么','核对','撑到'))})
  default=re.findall(r'(?:无(?:人)?处理|无人回应|未处理)([^\n]+)',text)
  cards[ident]={'id':ident,'title':title,'text':text,'eligibility':meta,'options':opts,'default':default[-1] if default else '', 'pool':'camp' if ident.startswith('C') else ('travel' if ident.startswith(('V','N-C')) else ('opportunity' if ident.startswith(('R','S')) else 'person'))}
 if re.match(r'^E\d{2} ',h) and b['style']=='Heading 3' and h[:3] not in endings and b['lines'] and b['lines'][0].startswith('判定：'):endings[h[:3]]={'title':h[4:],'text':'\n\n'.join(b['lines'])}
 if re.match(r'^S0[1-6] ',h) and b['style']=='Heading 3' and '合格' not in '\n'.join(b['lines'])[:90]:epilogues[h[:3]]={'title':h[4:],'text':'\n\n'.join(b['lines'])}
# The authoritative manuscript contains the original markdown tables as real Word tables.
mods={};locations=[];sources=[];slot='';days={}
for t in doc.tables:
 rows=[[c.text.strip() for c in row.cells] for row in t.rows]
 if rows and rows[0][0]=='来源场景':
  for cells in rows[1:]:
   sources.append({'id':'SRC%02d'%(len(sources)+1),'name':cells[0],'chance':int(re.search(r'\d+',cells[1])[0]),'choices':re.findall(r'SP\d{2}',cells[2]),'objects':cells[3]})
 if rows and rows[0][:2]==['日','职责']:
  for cells in rows[1:]:
   day=re.search(r'\d+',cells[0])
   if day:days[day[0]]={'title':cells[1],'scene':cells[2],'consequence':cells[3]}
 for cells in rows:
  if re.fullmatch(r'SP\d{2}',cells[0] if cells else ''):
   ident=cells[0];txt=' '.join(cells)
   # Full manuscript compact table: name/core, cost/weight, effect/limit.
   if len(cells)>=4 and '普通件' in cells[2]:
    n=re.findall(r'\d+',cells[2]);name=cells[1].split('核心')[0];core=cells[1].split('核心')[-1].strip()
    if len(n)>=3:mods[ident]={'id':ident,'name':name,'core':core,'cost':int(n[0]),'core_weight':int(n[1]),'weight':int(n[2]),'effect':cells[3]}
  if cells and re.fullmatch(r'LOC-SP\d{2}',cells[0]):
   locations.append({'id':cells[0],'name':cells[1],'chapter':int(cells[2]),'choices':re.findall(r'SP\d{2}',cells[3])})
# Resolve mechanical slots by identity, independent of table order.
slots={'front':[5,11,12,13,14,15],'engine':[1,10,16,17,18,19],'wheels':[2,6,20,21,22,23],'roof':[3,7,9,24,25,26],'rear':[4,8,27,28,29,30]}
for slot,nums in slots.items():
 for n in nums:
  if f'SP{n:02}' in mods:mods[f'SP{n:02}']['slot']=slot
story={};stages={};in_stage=False
for b in blocks:
 if b['heading']=='通用阶段与固定招募对白':in_stage=True
 elif b['style']=='Heading 1':in_stage=False
 if in_stage and b['style']=='Heading 2':stages[b['heading']]='\n\n'.join(b['lines'])
 if b['heading'].startswith('重要节点对白'):
  current=None
  for line in b['lines']:
   m=re.match(r'第(\d+)天[：:](.+)',line)
   if m:current=m[1];story[current]={'title':m[2],'lines':[]}
   elif current:story[current]['lines'].append(line)
raw={'source':docpath.name,'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'cards':cards,'endings':endings,'epilogues':epilogues,'modules':mods,'locations':locations,'sources':sources,'days':days,'story':story,'stages':stages}
(out/'content.json').write_text(json.dumps(raw,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'cards':len(cards),'endings':len(endings),'epilogues':len(epilogues),'modules':len(mods),'locations':len(locations),'missing':sorted(set(['M%02d'%i for i in range(1,16)]+['T%02d'%i for i in range(1,16)]+['V%02d'%i for i in range(1,12)]+['S%02d'%i for i in range(1,6)]+['R%02d'%i for i in range(1,9)]+['C%02d'%i for i in range(1,31)]+['N-F%02d'%i for i in range(1,12)]+['N-X%02d'%i for i in range(1,10)]+['N-C%02d'%i for i in range(1,6)])-set(cards))},ensure_ascii=False))

